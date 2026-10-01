#!/usr/bin/env python3
"""Fill Sheet1 keywords (column D) and categories (column E) with a local Ollama model.

Keywords are only ever taken from "Copy of Sheet1"; the model never writes new ones.

1. Downloads the latest workbook with sync_google_sheet.download_workbook_bytes()
   (or reads --workbook PATH).
2. Reads the real keyword lists in "Copy of Sheet1". Cells that only hold a slug
   ("a1-cassava") or the product name again are placeholders and are ignored.
3. A Sheet1 product whose name matches a Copy of Sheet1 item (score >= --reuse-threshold)
   copies that item's keywords as-is. Every other product is shown the keywords of its
   --similar closest Copy of Sheet1 items (nomic-embed-text embeddings of the names with
   brand words removed) and the model judges each one; only keywords it says fit are kept,
   and anything not in those lists is discarded.
4. The model also picks a category (from CATEGORIES) for every product. Answers are
   cached per product and model in reports/, so a run can be stopped with Ctrl+C and resumed.
5. Writes reports/sheet1_keywords_categories.csv for review and
   reports/sheet1_DE_paste.tsv to paste into Sheet1!D2, or with --push writes
   Sheet1!D2:E straight to Google Sheets using a service-account key.

Categories in "Copy of Sheet1" are not reused: nearly all are "Uncategorized" and the
rest are aisle names that often do not fit the product.

Usage:
  python3 scripts/fill_keywords_categories.py --limit 20        # quick trial
  python3 scripts/fill_keywords_categories.py                   # full run (resumable)
  python3 scripts/fill_keywords_categories.py --workbook store_workbook.xlsx \\
      --push --credentials service-account.json
"""

import argparse
import ast
import csv
import io
import json
import re
import sys
import time
from collections import Counter
from pathlib import Path
from urllib import error, request

import numpy as np
import openpyxl
from rapidfuzz import fuzz, process

sys.path.insert(0, str(Path(__file__).resolve().parent))
import sync_google_sheet  # noqa: E402

PROJECT_ROOT = Path(__file__).resolve().parent.parent
REPORTS_DIR = PROJECT_ROOT / "reports"
CACHE_PATH = REPORTS_DIR / "ollama_keyword_cache.json"
REVIEW_CSV_PATH = REPORTS_DIR / "sheet1_keywords_categories.csv"
PASTE_TSV_PATH = REPORTS_DIR / "sheet1_DE_paste.tsv"

OLLAMA_CHAT_URL = "http://localhost:11434/api/chat"
OLLAMA_EMBED_URL = "http://localhost:11434/api/embed"
DEFAULT_MODEL = "llama3.1:8b"
EMBED_MODEL = "nomic-embed-text:latest"
EMBED_PREFIX = "clustering: "  # nomic-embed-text task prefix for comparing texts to each other
CACHE_VERSION = "judge-v4"  # bump when the prompt or candidate rules change

TARGET_SHEET = "Sheet1"
REFERENCE_SHEET = "Copy of Sheet1"
# 1-based Sheet1 columns; sync_google_sheet.py reads keywords/categories from D/E.
COL_EN, COL_KEYWORDS, COL_CATEGORIES, COL_AISLE, COL_RACK = 1, 4, 5, 6, 7
EMPTY_CATEGORY_VALUES = {"", "uncategorized"}

CATEGORIES = [
    "Rice",
    "Atta & Flours",
    "Dals & Pulses",
    "Rava, Poha & Grains",
    "Spices & Masalas",
    "Dry Fruits & Nuts",
    "Oils & Ghee",
    "Dairy & Paneer",
    "Sugar, Salt & Jaggery",
    "Frozen Foods",
    "Snacks & Namkeen",
    "Sweets & Mithai",
    "Biscuits & Bakery",
    "Pickles & Chutneys",
    "Sauces & Pastes",
    "Ready to Eat & Instant",
    "Tea & Coffee",
    "Beverages",
    "Fruits & Vegetables",
    "Personal Care",
    "Household & Cleaning",
    "Pooja Items",
    "Kitchenware",
    "Other",
]

SYSTEM_PROMPT = f"""You label products for an Indian grocery store so shoppers can find them by search.
For each numbered product return:
- options: judge every option on that product's own options line, in order. The options are
  search keywords of other catalog items, so most do not fit.
  meaning: what the option means in plain English (translate Hindi and Telugu words).
  fits: true only if the option names this exact product's type, a synonym, or its Hindi or
  Telugu name. false if it names a different product, brand, flavour, variety, ingredient,
  dish or size, or is too vague to help find this product.
- category: exactly one label from the allowed list.
  Personal Care is for things used on the body (soap, shampoo, hair oil, toothpaste, cosmetics).
  Household & Cleaning is for the home (detergent, cleaners, gloves, foil, bags, disposables).
  Spices & Masalas includes whole spices, spice seeds (kalonji, methi, ajwain, mustard) and
  masala or seasoning mixes for a dish (tikka, biryani, chaat).
  Dry Fruits & Nuts includes dry and grated coconut.

Allowed categories: {"; ".join(CATEGORIES)}"""

RESPONSE_SCHEMA = {
    "type": "object",
    "properties": {
        "items": {
            "type": "array",
            "items": {
                "type": "object",
                "properties": {
                    "id": {"type": "integer"},
                    # the model sometimes invents an endless description here; real names are under 70 chars
                    "product": {"type": "string", "maxLength": 100},
                    "category": {"type": "string", "enum": CATEGORIES},
                    "options": {
                        "type": "array",
                        "items": {
                            "type": "object",
                            "properties": {
                                "option": {"type": "string"},
                                "meaning": {"type": "string"},
                                "fits": {"type": "boolean"},
                            },
                            "required": ["option", "meaning", "fits"],
                        },
                    },
                },
                "required": ["id", "product", "category", "options"],
            },
        }
    },
    "required": ["items"],
}

SIZE_RE = re.compile(
    r"\b\d+(?:\.\d+)?\s*(?:lbs?|kgs?|gms?|grams?|gr|g|oz|ml|ltrs?|litres?|liters?|l|"
    r"pcs?|pieces?|packs?|pk|ct|count|x)\b"
)
MAX_KEYWORD_WORDS = 4
MAX_KEYWORDS = 8
MAX_OPTIONS = 12  # keywords of similar items offered to the model per product
# A word that ends at least this many product names ("rice", "masala", "pickle") is a product type.
TYPE_WORD_MIN_COUNT = 5
# A word that starts at least this many names, and nearly always comes first, is a brand ("laxmi").
BRAND_WORD_MIN_COUNT = 3
BRAND_WORD_MIN_SHARE = 0.8


def singular(token):
    if len(token) > 3 and token.endswith("s") and not token.endswith("ss"):
        return token[:-1]
    return token


def normalize_name(name):
    """Lowercase, drop sizes/numbers/punctuation and plural 's' so near-duplicates compare equal."""
    text = str(name or "").lower().replace("&", " and ")
    text = re.sub(r"['’`]", "", text)
    text = SIZE_RE.sub(" ", text)
    text = re.sub(r"[^a-z0-9]+", " ", text)
    return " ".join(singular(t) for t in text.split() if not t.isdigit())


def name_tokens(name):
    return set(normalize_name(name).split())


def clean_keywords(parts):
    seen, cleaned = set(), []
    for part in parts:
        term = re.sub(r"[-_]+", " ", str(part).lower())
        term = re.sub(r"[\[\]'\"|]+", " ", term)  # a stray "|" would split the term in app.js
        term = re.sub(r"\s+", " ", term).strip(" .,;:")
        if term and re.search(r"[a-z]", term) and term not in seen:
            seen.add(term)
            cleaned.append(term)
    return cleaned


def parse_reference_keywords(raw):
    """Copy of Sheet1 mixes "a|b", "['a', 'b']", "a, b" and slug styles; return a clean list."""
    raw = (raw or "").strip()
    if not raw:
        return []
    if raw.startswith("[") and raw.endswith("]"):
        try:
            parsed = ast.literal_eval(raw)
            parts = [str(p) for p in parsed] if isinstance(parsed, (list, tuple)) else [raw]
        except (ValueError, SyntaxError):
            parts = re.findall(r"['\"]([^'\"]+)['\"]", raw) or [raw]
    elif "|" in raw:
        parts = raw.split("|")
    else:
        parts = raw.split(",")
    return clean_keywords(parts)


def reference_keywords(name, raw):
    """The real keyword list of a Copy of Sheet1 row, or [] for a slug or the name typed again."""
    terms = clean_keywords(SIZE_RE.sub(" ", term) for term in parse_reference_keywords(raw))
    terms = [term for term in terms if len(term.split()) <= MAX_KEYWORD_WORDS]
    if len(terms) == 1 and name_tokens(terms[0]) <= name_tokens(name):
        return []
    return [term for term in terms if normalize_name(term) != normalize_name(name)]


def product_type_words(names):
    counts = Counter(key.split()[-1] for key in map(normalize_name, names) if key)
    return {word for word, count in counts.items() if count >= TYPE_WORD_MIN_COUNT}


def brand_words(names, type_words):
    """Words that open many names and rarely appear anywhere else ("laxmi", "haldiram", "mdh")."""
    keys = [normalize_name(name).split() for name in names]
    first = Counter(key[0] for key in keys if len(key) > 1)
    anywhere = Counter(token for key in keys for token in set(key))
    brands = {word for word, count in first.items()
              if count >= BRAND_WORD_MIN_COUNT and count / anywhere[word] >= BRAND_WORD_MIN_SHARE}
    return brands - type_words


def strip_brands(name, brands):
    """The name without brand words, so "Laxmi Chilli Powder" is compared as "chilli powder"."""
    tokens = normalize_name(name).split()
    return " ".join([token for token in tokens if token not in brands] or tokens)


def candidate_terms(reference, product_name, type_words):
    """Keywords of a similar item that could also describe `product_name`.

    Terms already spelled out by the product's own name add nothing to search. A term cut
    from the similar item's own name ("24 mantra", "aachi pickle") describes that item, so it
    is only offered when its extra words are product types ("muruku" for "Murukulu");
    synonyms and translations ("dhania", "biyyam") are always offered.
    """
    target, source = name_tokens(product_name), name_tokens(reference["name"])
    terms = []
    for term in reference["keywords"]:
        tokens = name_tokens(term)
        if tokens <= target or (tokens <= source and not (tokens - target) <= type_words):
            continue
        terms.append(term)
    return terms


def cell(row, col):
    value = row[col - 1] if len(row) >= col else None
    if value is None:
        return ""
    if isinstance(value, float) and value.is_integer():
        value = int(value)
    return str(value).strip()


def load_workbook(path):
    if path:
        data = Path(path).read_bytes()
    else:
        data = sync_google_sheet.download_workbook_bytes()
        if not data:
            raise SystemExit("Could not download the workbook and no cached copy exists.")
    workbook = openpyxl.load_workbook(io.BytesIO(data), read_only=True, data_only=True)
    for name in (TARGET_SHEET, REFERENCE_SHEET):
        if name not in workbook.sheetnames:
            raise SystemExit(f'Workbook has no "{name}" sheet (found: {workbook.sheetnames}).')
    return workbook


def read_target_rows(workbook):
    rows = []
    for row_number, row in enumerate(workbook[TARGET_SHEET].iter_rows(min_row=2, values_only=True), start=2):
        rows.append({
            "row": row_number,
            "en": cell(row, COL_EN),
            "keywords": cell(row, COL_KEYWORDS),
            "categories": cell(row, COL_CATEGORIES),
            "aisle": cell(row, COL_AISLE),
            "rack": cell(row, COL_RACK),
        })
    while rows and not any(rows[-1][k] for k in ("en", "keywords", "categories", "aisle", "rack")):
        rows.pop()
    return rows


def read_references(workbook):
    """Return ({normalized name: {"name", "keywords"}} for real keyword lists, all reference names)."""
    references, names = {}, []
    for row in workbook[REFERENCE_SHEET].iter_rows(min_row=2, values_only=True):
        name = cell(row, COL_EN)
        key = normalize_name(name)
        if not key:
            continue
        names.append(name)
        keywords = reference_keywords(name, cell(row, COL_KEYWORDS))
        if keywords and (key not in references or len(keywords) > len(references[key]["keywords"])):
            references[key] = {"name": name, "keywords": keywords}
    return references, names


def group_products(rows, references, reuse_threshold):
    """One entry per normalized product name, with the Copy of Sheet1 item to copy keywords from."""
    products = {}
    for row in rows:
        key = normalize_name(row["en"])
        if key:
            products.setdefault(key, {"key": key, "name": row["en"], "rows": []})["rows"].append(row)

    reference_keys = list(references)
    for product in products.values():
        product["copy"], product["score"], product["similar"] = None, 0, []
        if not reference_keys:
            continue
        match_key, score, _ = process.extractOne(product["key"], reference_keys, scorer=fuzz.token_sort_ratio)
        if score >= reuse_threshold:
            product["copy"], product["score"] = references[match_key], round(score)
    return products


def load_cache():
    if CACHE_PATH.exists():
        return json.loads(CACHE_PATH.read_text(encoding="utf-8"))
    return {}


def save_cache(cache):
    REPORTS_DIR.mkdir(exist_ok=True)
    tmp_path = CACHE_PATH.with_suffix(".tmp")
    tmp_path.write_text(json.dumps(cache, ensure_ascii=False, indent=1), encoding="utf-8")
    tmp_path.replace(CACHE_PATH)


def post_json(url, payload, timeout):
    req = request.Request(url, data=json.dumps(payload).encode("utf-8"), headers={"Content-Type": "application/json"})
    with request.urlopen(req, timeout=timeout) as resp:
        return json.loads(resp.read().decode("utf-8"))


def ollama_embed(texts, timeout, batch_size=256):
    """Unit-length nomic-embed-text vectors, one row per text."""
    vectors = []
    for start in range(0, len(texts), batch_size):
        batch = [EMBED_PREFIX + text for text in texts[start : start + batch_size]]
        vectors.extend(post_json(OLLAMA_EMBED_URL, {"model": EMBED_MODEL, "input": batch}, timeout)["embeddings"])
    matrix = np.array(vectors, dtype=np.float32)
    return matrix / np.linalg.norm(matrix, axis=1, keepdims=True)


def attach_similar_items(products, references, type_words, brands, count, min_similarity, timeout):
    """Give each product up to `count` of the closest Copy of Sheet1 items that have terms to offer.

    Names are compared without brand words; otherwise every Laxmi product looks most like
    other Laxmi products.
    """
    products = [p for p in products if not p["copy"]]
    if not products or not references or count <= 0:
        return
    print(f"🔎 Finding similar {REFERENCE_SHEET} items with {EMBED_MODEL}...")
    reference_list = list(references.values())
    reference_vectors = ollama_embed([strip_brands(r["name"], brands) for r in reference_list], timeout)
    product_vectors = ollama_embed([strip_brands(p["name"], brands) for p in products], timeout)
    for product, similarities in zip(products, product_vectors @ reference_vectors.T):
        product["similar"] = []
        for index in np.argsort(-similarities)[: count * 3]:
            if similarities[index] < min_similarity:
                break
            reference = reference_list[index]
            terms = candidate_terms(reference, product["name"], type_words)
            if terms:
                product["similar"].append({"name": reference["name"], "keywords": terms, "score": float(similarities[index])})
            if len(product["similar"]) == count:
                break


def ollama_chat(model, user_prompt, timeout):
    payload = {
        "model": model,
        "stream": False,
        "format": RESPONSE_SCHEMA,
        "keep_alive": "30m",
        "options": {"temperature": 0, "num_ctx": 8192, "num_predict": 4096},  # stop runaway answers early
        "messages": [
            {"role": "system", "content": SYSTEM_PROMPT},
            {"role": "user", "content": user_prompt},
        ],
    }
    body = post_json(OLLAMA_CHAT_URL, payload, timeout)
    return json.loads(body["message"]["content"])


def keyword_pool(product):
    """{option: similar item it came from}, closest similar item first."""
    pool = {}
    for similar in product["similar"]:
        for term in similar["keywords"]:
            if len(pool) < MAX_OPTIONS:
                pool.setdefault(term, similar["name"])
    return pool


def build_prompt(batch):
    lines = []
    for index, product in enumerate(batch, start=1):
        options = " | ".join(keyword_pool(product)) or "(none)"
        lines.append(f"{index}. {product['name']}\n   options: {options}")
    lines.append(f"Return exactly {len(batch)} items: id is the product number, product is its name.")
    return "\n".join(lines)


def label_batch(model, batch, timeout):
    """Return {product key: {"category", "keywords", "from"}}; products the model skipped are retried alone."""
    try:
        items = ollama_chat(model, build_prompt(batch), timeout).get("items", [])
    except (error.URLError, TimeoutError, ValueError, KeyError) as exc:
        if len(batch) == 1:
            print(f"   ⚠️ Ollama failed for {batch[0]['name']!r}: {exc}")
            return {}
        items = []

    labels = {}
    for item in items:
        index = item.get("id")
        if not isinstance(index, int) or not 1 <= index <= len(batch) or item.get("category") not in CATEGORIES:
            continue
        product = batch[index - 1]
        if fuzz.token_set_ratio(normalize_name(item.get("product")), product["key"]) < 80:
            continue  # answer belongs to another product; the product is retried on its own
        pool = keyword_pool(product)
        fitting = [option.get("option") for option in item.get("options") or [] if option.get("fits") is True]
        picked = [term for term in clean_keywords(fitting) if term in pool][:MAX_KEYWORDS]
        labels.setdefault(product["key"], {
            "category": item["category"],
            "keywords": picked,
            "from": list(dict.fromkeys(pool[term] for term in picked)),
        })

    if len(batch) > 1:
        for product in batch:
            if product["key"] not in labels:
                labels.update(label_batch(model, [product], timeout))
    return labels


def needs_model(product, overwrite):
    """True if a row lacks a category, or lacks keywords that cannot be copied as-is."""
    missing_category = overwrite or any(row["categories"].lower() in EMPTY_CATEGORY_VALUES for row in product["rows"])
    missing_keywords = overwrite or any(not row["keywords"] for row in product["rows"])
    return missing_category or (missing_keywords and not product["copy"])


def run_ollama(products, references, type_words, brands, cache, args):
    model_cache = cache.setdefault(f"{args.model}|{CACHE_VERSION}", {})
    pending = sorted(
        (p for p in products.values() if p["key"] not in model_cache and needs_model(p, args.overwrite)),
        key=lambda p: p["rows"][0]["row"],
    )
    if args.limit is not None:
        pending = pending[: args.limit]
    if not pending:
        return
    attach_similar_items(pending, references, type_words, brands, args.similar, args.min_similarity, args.timeout)
    print(f"🤖 Asking {args.model} about {len(pending)} products in batches of {args.batch_size}...", flush=True)

    started = time.time()
    try:
        for start in range(0, len(pending), args.batch_size):
            batch = pending[start : start + args.batch_size]
            model_cache.update(label_batch(args.model, batch, args.timeout))
            save_cache(cache)
            done = start + len(batch)
            rate = done / max(time.time() - started, 1e-6)
            eta_min = (len(pending) - done) / rate / 60
            print(f"   [{done}/{len(pending)}] {rate * 60:.1f} products/min, ~{eta_min:.0f} min left", flush=True)
    except KeyboardInterrupt:
        print("\n⏸️ Interrupted; progress is cached. Re-run the same command to resume.")


# The model is inconsistent on body-care items it can also read as food/cleaning:
# hair oil lands in "Oils & Ghee" and bath soap in "Household & Cleaning" about half
# the time, though SYSTEM_PROMPT says both belong in Personal Care. These names are
# unambiguous, so the rule decides them instead of the model. "Dish soap" and laundry
# soap are genuinely Household & Cleaning and are left alone.
PERSONAL_CARE_PATTERNS = re.compile(
    r"hair oil|hair colou?r|shampoo|conditioner|toothpaste|tooth paste|"
    r"hand ?wash|face wash|body lotion|body wash|talc|deodorant|\bsoap\b",
    re.IGNORECASE,
)
NOT_PERSONAL_CARE_PATTERNS = re.compile(
    r"dish|dishwash|utensil|laundry|detergent|floor|toilet", re.IGNORECASE
)


def forced_category(name):
    """Return the category a product name settles on its own, else None."""
    if PERSONAL_CARE_PATTERNS.search(name) and not NOT_PERSONAL_CARE_PATTERNS.search(name):
        return "Personal Care"
    return None


def resolve_rows(rows, products, cache, args):
    model_cache = cache.get(f"{args.model}|{CACHE_VERSION}", {})
    resolved = []
    for row in rows:
        key = normalize_name(row["en"])
        product = products.get(key)
        label = model_cache.get(key) or {}
        copy = product["copy"] if product else None

        keywords, keyword_source, keywords_from = row["keywords"], "existing", ""
        if args.overwrite or not keywords:
            if copy:
                keywords, keyword_source, keywords_from = "|".join(copy["keywords"][:MAX_KEYWORDS]), "copy_of_sheet1", copy["name"]
            elif label.get("keywords"):
                keywords, keyword_source = "|".join(label["keywords"]), "picked_from_similar"
                keywords_from = "; ".join(label.get("from", []))
            else:
                keywords, keyword_source = "", ""

        category, category_source = row["categories"], "existing"
        if args.overwrite or category.lower() in EMPTY_CATEGORY_VALUES:
            category = label.get("category", "")
            category_source = "ollama" if category else ""
            forced = forced_category(row["en"])
            if forced and category != forced:
                category, category_source = forced, "rule"

        resolved.append({
            **row,
            "keywords": keywords,
            "categories": category,
            "keywords_source": keyword_source,
            "categories_source": category_source,
            "copy_match_score": product["score"] if copy else "",
            "keywords_from": keywords_from,
        })
    return resolved


def write_outputs(resolved):
    REPORTS_DIR.mkdir(exist_ok=True)
    fields = ["row", "en", "aisle", "rack", "keywords", "categories",
              "keywords_source", "categories_source", "copy_match_score", "keywords_from"]
    with REVIEW_CSV_PATH.open("w", encoding="utf-8", newline="") as f:
        writer = csv.DictWriter(f, fieldnames=fields, extrasaction="ignore")
        writer.writeheader()
        writer.writerows(resolved)
    with PASTE_TSV_PATH.open("w", encoding="utf-8", newline="") as f:
        for row in resolved:
            f.write(f"{row['keywords']}\t{row['categories']}\n")


def push_to_sheet(resolved, credentials_path):
    """Write Sheet1!D2:E after checking column A still matches the downloaded copy row for row."""
    from google.oauth2 import service_account
    from googleapiclient.discovery import build

    creds = service_account.Credentials.from_service_account_file(
        credentials_path, scopes=["https://www.googleapis.com/auth/spreadsheets"]
    )
    values_api = build("sheets", "v4", credentials=creds, cache_discovery=False).spreadsheets().values()
    spreadsheet_id = sync_google_sheet.SPREADSHEET_ID
    last_row = resolved[-1]["row"]

    live = values_api.get(spreadsheetId=spreadsheet_id, range=f"{TARGET_SHEET}!A2:A{last_row}").execute()
    live_names = [(r[0].strip() if r else "") for r in live.get("values", [])]
    live_names += [""] * (len(resolved) - len(live_names))
    for row, live_name in zip(resolved, live_names):
        if row["en"] != live_name:
            raise SystemExit(
                f"❌ {TARGET_SHEET} changed since download (row {row['row']}: "
                f"{row['en']!r} vs {live_name!r}). Re-run without --workbook to refresh."
            )

    values_api.update(
        spreadsheetId=spreadsheet_id,
        range=f"{TARGET_SHEET}!D2:E{last_row}",
        valueInputOption="RAW",
        body={"values": [[row["keywords"], row["categories"]] for row in resolved]},
    ).execute()
    print(f"☁️ Wrote {TARGET_SHEET}!D2:E{last_row} to Google Sheets.")


def main():
    parser = argparse.ArgumentParser(description=__doc__, formatter_class=argparse.RawDescriptionHelpFormatter)
    parser.add_argument("--workbook", help="read this .xlsx instead of downloading the latest copy")
    parser.add_argument("--model", default=DEFAULT_MODEL, help=f"Ollama model (default {DEFAULT_MODEL})")
    parser.add_argument("--batch-size", type=int, default=5, help="products per Ollama request")
    parser.add_argument("--limit", type=int, help="only ask the model about this many new products (0 = none)")
    parser.add_argument("--timeout", type=int, default=600, help="seconds per Ollama request")
    parser.add_argument("--reuse-threshold", type=int, default=95,
                        help="name match score (0-100) at which Copy of Sheet1 keywords are copied as-is")
    parser.add_argument("--similar", type=int, default=5,
                        help="similar Copy of Sheet1 items whose keywords the model may pick from")
    parser.add_argument("--min-similarity", type=float, default=0.8,
                        help="cosine similarity (brand words removed) below which an item is not similar")
    parser.add_argument("--overwrite", action="store_true", help="replace keywords/categories already in Sheet1")
    parser.add_argument("--push", action="store_true", help="write Sheet1!D:E back to Google Sheets")
    parser.add_argument("--credentials", help="service-account JSON key with edit access to the sheet (--push)")
    args = parser.parse_args()
    if args.push and not args.credentials:
        parser.error("--push needs --credentials")

    workbook = load_workbook(args.workbook)
    rows = read_target_rows(workbook)
    references, reference_names = read_references(workbook)
    all_names = [row["en"] for row in rows] + reference_names
    type_words = product_type_words(all_names)
    brands = brand_words(all_names, type_words)
    products = group_products(rows, references, args.reuse_threshold)
    copied = sum(1 for p in products.values() if p["copy"])
    print(f"📋 {TARGET_SHEET}: {len(rows)} rows, {len(products)} distinct products. "
          f"{REFERENCE_SHEET}: {len(references)} of {len(reference_names)} items have real keyword lists.")
    print(f"🔁 {copied} products copy keywords as-is; the rest pick from their "
          f"{args.similar} most similar {REFERENCE_SHEET} items.", flush=True)

    cache = load_cache()
    run_ollama(products, references, type_words, brands, cache, args)

    resolved = resolve_rows(rows, products, cache, args)
    write_outputs(resolved)
    named = [r for r in resolved if r["en"]]
    missing_keywords = sum(1 for r in named if not r["keywords"])
    missing_categories = sum(1 for r in named if not r["categories"])
    print(f"📝 Review: {REVIEW_CSV_PATH.relative_to(PROJECT_ROOT)}")
    print(f"📋 Paste into {TARGET_SHEET}!D2: {PASTE_TSV_PATH.relative_to(PROJECT_ROOT)}")
    print(f"   Rows still missing keywords: {missing_keywords}, categories: {missing_categories}")

    if args.push:
        if missing_keywords or missing_categories:
            print("⚠️ Pushing a partial fill; re-run later to complete the remaining rows.")
        push_to_sheet(resolved, args.credentials)


if __name__ == "__main__":
    main()
